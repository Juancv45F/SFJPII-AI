import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from services.cliente_tools import CLIENTE_TOOLS, ejecutar_cliente_tool
from tools import informacion_farmacia
from tools.informacion_farmacia import consultar_informacion_farmacia


class ClientPharmacyInfoTest(unittest.TestCase):
    def test_production_document_contains_current_contact_and_schedule(self):
        result = consultar_informacion_farmacia()

        self.assertTrue(result["disponible"])
        self.assertIn("Información de la farmacia", result["contenido"])
        self.assertIn("Prolongación Cristóbal Colón", result["contenido"])
        self.assertIn("495 127 8814", result["contenido"])
        self.assertIn("9:00 a. m. a 9:00 p. m.", result["contenido"])
        self.assertIn("Excepciones de horario vigentes", result["contenido"])

    def test_reads_a_saved_document_change_without_restarting(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "farmacia.docx"
            document = Document()
            document.add_paragraph("Horario: 9:00 a. m. a 9:00 p. m.")
            document.save(path)

            with patch.object(
                informacion_farmacia.settings,
                "pharmacy_info_docx",
                str(path),
                create=True,
            ):
                first = consultar_informacion_farmacia()
                document = Document()
                document.add_paragraph("Horario especial: Cerrado el 25 de diciembre de 2026")
                document.save(path)
                second = consultar_informacion_farmacia()

        self.assertIn("9:00 a. m.", first["contenido"])
        self.assertIn("Cerrado el 25 de diciembre de 2026", second["contenido"])
        self.assertNotIn("9:00 a. m.", second["contenido"])

    def test_tool_has_no_file_path_argument_and_is_client_allowed(self):
        schema = next(
            tool for tool in CLIENTE_TOOLS
            if tool["name"] == "consultar_informacion_farmacia"
        )
        self.assertEqual(schema["parameters"]["properties"], {})

        result = ejecutar_cliente_tool(
            "consultar_informacion_farmacia", {}, user_id="ignored-by-public-tool"
        )
        self.assertTrue(result["disponible"])


if __name__ == "__main__":
    unittest.main()
