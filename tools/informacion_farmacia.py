from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from config import settings


PROJECT_ROOT = Path(__file__).resolve().parent.parent
MAX_DOCUMENT_CHARS = 20_000


def _document_path() -> Path:
    configured = Path(
        getattr(settings, "pharmacy_info_docx", "data/informacion_farmacia.docx")
    ).expanduser()
    if not configured.is_absolute():
        configured = PROJECT_ROOT / configured
    return configured.resolve()


def _read_document(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if text:
            blocks.append(text)

    for index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            values = [" ".join(cell.text.split()) for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            blocks.append(f"Tabla {index}:\n" + "\n".join(rows))

    return "\n\n".join(blocks).strip()


def consultar_informacion_farmacia():
    """Lee el documento configurado; el modelo nunca recibe ni decide su ruta."""
    path = _document_path()
    if not path.is_file():
        return {
            "disponible": False,
            "mensaje": "El documento de información de la farmacia no está disponible."
        }

    try:
        content = _read_document(path)
    except (OSError, ValueError, KeyError, PackageNotFoundError):
        return {
            "disponible": False,
            "mensaje": "No fue posible leer el documento de información de la farmacia."
        }

    if not content:
        return {
            "disponible": False,
            "mensaje": "El documento de información de la farmacia está vacío."
        }

    if len(content) > MAX_DOCUMENT_CHARS:
        content = content[:MAX_DOCUMENT_CHARS].rstrip() + "\n[Contenido recortado]"

    updated = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
    return {
        "disponible": True,
        "actualizado_en": updated.isoformat(),
        "contenido": content,
    }
